# Create an Excel template for PERIODE_DATA using openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

wb = Workbook()

# Main data sheet
ws_main = wb.active
ws_main.title = "Lampiran"

headers = [
    "No",
    "Nama",
    "Email",
    "ID Peserta",
    "Periode",
    "Status Reservasi",
    "Status Lulus",
    "Tanggal Update",
    "Updated By"
]

ws_main.append(headers)

for cell in ws_main[1]:
    cell.font = Font(bold=True)
    cell.alignment = Alignment(horizontal="center")

for i in range(len(headers)):
    ws_main.column_dimensions[get_column_letter(i+1)].width = 20

# Add sample rows
sample_rows = [
    [1, "Rina Kurnia", "rina@email.com", "PRD001", "2026-A", "Sudah", "Lulus", "2026-03-01", "Admin"],
    [2, "Budi Santoso", "budi@email.com", "PRD002", "2026-A", "Belum", "Belum", "2026-03-02", "Admin"],
    [3, "Lina Wijaya", "lina@email.com", "PRD003", "2026-B", "Sudah", "Belum", "2026-03-03", "Validator"],
]

for r in sample_rows:
    ws_main.append(r)

# Rekapan sheet
ws_rekap = wb.create_sheet("Rekapan")

ws_rekap.append(["Kategori", "Jumlah"])

rows = [
    ["Total Data", '=COUNTA(Lampiran!A:A)-1'],
    ["Reservasi", '=COUNTIF(Lampiran!F:F,"Sudah")'],
    ["Belum Reservasi", '=COUNTIF(Lampiran!F:F,"Belum")'],
    ["Lulus", '=COUNTIF(Lampiran!G:G,"Lulus")'],
    ["Belum Lulus", '=COUNTIF(Lampiran!G:G,"Belum")'],
]

for r in rows:
    ws_rekap.append(r)

for cell in ws_rekap[1]:
    cell.font = Font(bold=True)

ws_rekap.column_dimensions["A"].width = 30
ws_rekap.column_dimensions["B"].width = 20

# Periode_List sheet (reference for periods)
ws_periode = wb.create_sheet("Periode_List")
ws_periode.append(["Kode Periode", "Nama Periode", "Tahun", "Keterangan"])

for cell in ws_periode[1]:
    cell.font = Font(bold=True)

periode_rows = [
    ["2026-A", "Periode Awal 2026", "2026", "Batch pertama"],
    ["2026-B", "Periode Tengah 2026", "2026", "Batch kedua"],
]

for r in periode_rows:
    ws_periode.append(r)

for i in range(4):
    ws_periode.column_dimensions[get_column_letter(i+1)].width = 25

# Print view sheet
ws_print = wb.create_sheet("Print_View")
ws_print.append(headers)

for cell in ws_print[1]:
    cell.font = Font(bold=True)

for i in range(len(headers)):
    ws_print.column_dimensions[get_column_letter(i+1)].width = 20

path = "PERIODE_DATA_Template.xlsx"
wb.save(path)

path